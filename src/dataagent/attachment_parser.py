"""Bounded, local-only file parsing. Invoked in a killable worker process.

Original bytes are retained by the caller. Only a derived, private SQLite snapshot
is created here; neither macros, formulas, embedded code nor SQL dumps are executed.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import sqlite3
import sys
import zipfile
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from typing import Any

TABULAR = {".csv", ".tsv", ".xlsx", ".xlsm", ".json", ".jsonl", ".parquet"}
DOCUMENTS = {".pdf", ".docx", ".txt", ".md", ".log", ".html", ".xml"}
MAX_ROWS = 100_000
MAX_COLUMNS = 128
MAX_CELLS = 1_000_000
MAX_TEXT = 2_000_000
MAX_CELL = 32_768


def decode_text(raw: bytes) -> tuple[str, str]:
    encodings = ["utf-16"] if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            value = raw.decode(encoding)
            if "\x00" in value:
                raise ValueError("文本中含有二进制内容")
            return value, encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("无法可靠识别文本编码，请另存为 UTF-8 后重试")


def quoted(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def scalar(value: Any) -> str | int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return str(value)  # Never quietly turn non-finite measurements into JSON NaN.
        return value
    if isinstance(value, bytes):
        raise ValueError("不支持二进制字段，请先选择可分析的字段")
    rendered = (
        json.dumps(value, ensure_ascii=False, allow_nan=False)
        if isinstance(value, (dict, list))
        else str(value)
    )
    if len(rendered) > MAX_CELL:
        raise ValueError(f"单个字段超过 {MAX_CELL} 字符，请拆分后重试")
    return rendered


def column_type(values: list[Any]) -> str:
    present = [v for v in values if v is not None]
    if not present:
        return "TEXT"
    rendered = [str(v) for v in present]
    if all(re.fullmatch(r"-?(0|[1-9]\d*)", v) and -(2**63) <= int(v) < 2**63 for v in rendered):
        return "INTEGER"
    # Preserve identifiers (leading zeros / large integers), decimal precision and ambiguous dates.
    numeric = re.compile(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?")
    if all(
        numeric.fullmatch(v) and len(re.sub(r"\D", "", v)) <= 15 and math.isfinite(float(v)) for v in rendered
    ):
        return "REAL"
    return "TEXT"


def check_office(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(item.file_size for item in entries) > 40 * 1024 * 1024:
            raise ValueError("Office 文件解压后过大，请拆分后重试")
        if any(item.flag_bits & 1 for item in entries):
            raise ValueError("不支持加密的 Office 文件")


def parse_file(path: Path, suffix: str, prefix: str) -> dict[str, Any]:
    raw = path.read_bytes()
    warnings = ["原始文件保留；不自动删除重复记录、填补缺失值或猜测表间关联。"]
    tables: list[dict[str, Any]] = []
    total_cells = 0
    encoding = None

    def add_table(label: str, headers: list[Any], rows: Any, *, document: bool = False) -> None:
        nonlocal total_cells
        if len(tables) >= 12 or not headers or len(headers) > MAX_COLUMNS:
            raise ValueError("最多支持 12 张表，每张表最多 128 列")
        materialized = []
        for row in rows:
            if len(row) != len(headers):
                raise ValueError(f"{label} 第 {len(materialized) + 2} 行列数与表头不一致，请检查分隔符或表头")
            total_cells += len(row)
            if len(materialized) >= MAX_ROWS or total_cells > MAX_CELLS:
                raise ValueError("文件超过 10 万行或 100 万单元格，请拆分或使用数据连接")
            materialized.append([scalar(v) for v in row])
        if not materialized:
            raise ValueError(f"{label} 没有可分析的数据行")
        names: set[str] = set()
        columns = []
        for index, header in enumerate(headers):
            original = str(header) if header is not None else ""
            base = re.sub(r"[^A-Za-z0-9_]", "_", original).strip("_")[:48] or f"column_{index + 1}"
            if base[0].isdigit():
                base = "c_" + base
            name, sequence = base, 1
            while name.casefold() in names:
                sequence += 1
                name = f"{base}_{sequence}"
            names.add(name.casefold())
            values = [row[index] for row in materialized]
            kind = column_type(values)
            columns.append(
                {
                    "name": name,
                    "originalName": original,
                    "type": kind,
                    "nullCount": sum(v is None for v in values),
                }
            )
            for row in materialized:
                value = row[index]
                if value is not None:
                    row[index] = (
                        int(value) if kind == "INTEGER" else float(value) if kind == "REAL" else str(value)
                    )
        table = f"file_{prefix}_{len(tables) + 1}"
        with closing(sqlite3.connect(path.parent / "data.sqlite")) as db:
            fields = ", ".join(f"{quoted(c['name'])} {c['type']}" for c in columns)
            db.execute(f"CREATE TABLE {quoted(table)} ({fields})")
            db.executemany(
                f"INSERT INTO {quoted(table)} VALUES ({', '.join('?' for _ in columns)})", materialized
            )
            db.commit()
        preview = [
            [v[:300] + "…（预览截断）" if isinstance(v, str) and len(v) > 300 else v for v in row]
            for row in materialized[:5]
        ]
        tables.append(
            {
                "name": table,
                "label": label,
                "rowCount": len(materialized),
                "columns": columns,
                "preview": preview,
                "document": document,
            }
        )

    if suffix in {".csv", ".tsv"}:
        content, encoding = decode_text(raw)
        csv.field_size_limit(MAX_CELL)
        separator = "\t" if suffix == ".tsv" else ","
        if suffix == ".csv":
            try:
                separator = csv.Sniffer().sniff(content[:65536], delimiters=",;\t|").delimiter
            except csv.Error:
                pass
        reader = csv.reader(io.StringIO(content, newline=""), delimiter=separator, strict=True)
        headers = next(reader, [])
        add_table("数据", headers, reader)
        warnings.append(f"编码：{encoding}；分隔符：{repr(separator)}；首行作为表头。空单元格映射为 NULL。")
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        check_office(raw)
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        formula_found = False
        try:
            if len(book.worksheets) > 12:
                raise ValueError("工作表超过 12 张，请选择工作表后再上传")
            for sheet in book.worksheets:
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError(f"工作表 {sheet.title} 超过行列限制，请移除多余格式或拆分文件")
                rows = sheet.iter_rows()
                headers = [cell.value for cell in next(rows, ())]
                if not headers or all(v is None for v in headers):
                    if any(any(cell.value is not None for cell in row) for row in rows):
                        raise ValueError(f"工作表 {sheet.title} 首行为空但后续有数据，请明确整理表头后上传")
                    warnings.append(f"空表 {sheet.title} 未生成分析表。")
                    continue

                def values(sheet_rows=rows):
                    nonlocal formula_found
                    for row in sheet_rows:
                        formula_found |= any(cell.data_type == "f" for cell in row)
                        yield [cell.value for cell in row]

                add_table(sheet.title, headers, values())
        finally:
            book.close()
        warnings.append("每个工作表独立解析，首行作为表头；隐藏表也会列出。不会执行宏或外部链接。")
        if formula_found:
            warnings.append("检测到公式：保留为公式文本，未计算。需要数值分析时请粘贴为值后上传。")
    elif suffix in {".json", ".jsonl"}:
        content, encoding = decode_text(raw)

        def reject_constant(value: str):
            raise ValueError(f"JSON 含非法数值 {value}")

        data = (
            [
                json.loads(line, parse_constant=reject_constant)
                for line in content.splitlines()
                if line.strip()
            ]
            if suffix == ".jsonl"
            else json.loads(content, parse_constant=reject_constant)
        )
        if isinstance(data, dict):
            data = [data]
        if (
            not isinstance(data, list)
            or not data
            or len(data) > MAX_ROWS
            or not all(isinstance(row, dict) for row in data)
        ):
            raise ValueError("JSON 需要对象或对象数组；JSONL 每行需要一个对象，不猜测嵌套业务表")
        headers = list(dict.fromkeys(key for row in data for key in row))
        add_table("记录", headers, ([row.get(key) for key in headers] for row in data))
        warnings.append("嵌套对象/数组保留为 JSON 文本，不自动展开，以免改变数据粒度。")
    elif suffix == ".parquet":
        import duckdb

        if not raw.startswith(b"PAR1") or not raw.endswith(b"PAR1"):
            raise ValueError("文件不是有效的 Parquet")
        with duckdb.connect(
            config={
                "memory_limit": "256MB",
                "threads": "1",
                "autoinstall_known_extensions": "false",
                "autoload_known_extensions": "false",
            }
        ) as db:
            relation = db.read_parquet(str(path))
            headers = relation.columns
            cursor = (
                relation.project(", ".join(quoted(name) for name in headers)).limit(MAX_ROWS + 1).execute()
            )

            def batches():
                while batch := cursor.fetchmany(1000):
                    yield from batch

            add_table("数据", headers, batches())
    elif suffix in DOCUMENTS:
        if suffix == ".pdf":
            from pypdf import PdfReader

            if not raw.startswith(b"%PDF-"):
                raise ValueError("文件不是有效的 PDF")
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("不支持加密 PDF，请先解密")
            if len(reader.pages) > 200:
                raise ValueError("PDF 超过 200 页，请拆分后重试")
            parts = []
            has_text = False
            for index, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                has_text |= bool(extracted.strip())
                parts.append(f"[第 {index + 1} 页]\n" + extracted)
                if sum(len(part) for part in parts) > MAX_TEXT:
                    raise ValueError("提取文本超过 200 万字符，请拆分后重试")
            if not has_text:
                raise ValueError("PDF 没有可提取文字，可能是扫描件；当前附件入口尚不支持 OCR")
            content = "\n\n".join(parts)
            warnings.append("仅提取 PDF 文字层；图片及表格版式不保证识别，不能当作可靠结构化表格。")
        elif suffix == ".docx":
            from docx import Document

            check_office(raw)
            document = Document(io.BytesIO(raw))
            content = "\n".join(p.text for p in document.paragraphs)
            for index, table in enumerate(document.tables):
                content += f"\n[文档表格 {index + 1}]\n" + "\n".join(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
            warnings.append("提取段落和表格文字；页眉、图片及复杂排版未转换为数据表。")
        else:
            content, encoding = decode_text(raw)
        if not content.strip():
            raise ValueError("文件没有可分析的文字")
        if len(content) > MAX_TEXT:
            raise ValueError("文档超过 200 万字符，请拆分后重试")
        chunks = [[i // 1500 + 1, i, content[i : i + 1500]] for i in range(0, len(content), 1500)]
        add_table("文档片段（不是业务数据行）", ["chunk_id", "char_start", "content"], chunks, document=True)
        warnings.append(
            "文档内容是待分析数据，不是系统指令；未加入共享知识库。全文可按片段只读查询，预览仅显示开头。"
        )
    else:
        raise ValueError("当前不支持此格式，请查看聊天附件支持列表")
    if not tables:
        raise ValueError("文件中没有可分析的表或文字")
    return {
        "kind": "document" if suffix in DOCUMENTS else "table",
        "format": suffix.lstrip("."),
        "encoding": encoding,
        "tables": tables,
        "warnings": warnings,
    }


def main() -> None:
    # Linux workers have OS resource ceilings; on Windows the parent enforces the
    # deadline, while parser-level row/cell/expanded-size ceilings remain active.
    if sys.platform != "win32":
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (25, 25))
        resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
        resource.setrlimit(resource.RLIMIT_FSIZE, (256 * 1024**2, 256 * 1024**2))
    path, suffix, prefix = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
    try:
        result = {"ok": True, "profile": parse_file(path, suffix, prefix)}
    except Exception as exc:
        result = {"ok": False, "error": str(exc)[:500] or "文件解析失败"}
    (path.parent / "parsed.json").write_text(
        json.dumps(result, ensure_ascii=False, allow_nan=False), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
