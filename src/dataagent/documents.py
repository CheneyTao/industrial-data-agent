from __future__ import annotations

import asyncio
import csv
import io
import json
import re
from pathlib import Path
from typing import Protocol
from urllib.parse import urlparse
from uuid import uuid4

import httpx
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from dataagent.config import Settings, get_settings
from dataagent.errors import ConfigurationError, DataAgentError
from dataagent.llm import LLMClient, get_llm_client
from dataagent.models import AgentKnowledge, BusinessKnowledge
from dataagent.prompts import render_prompt
from dataagent.vector_metadata import (
    AGENT_ID,
    AGENT_KNOWLEDGE,
    AGENT_KNOWLEDGE_ID,
    BUSINESS_TERM,
    BUSINESS_TERM_ID,
    CONCRETE_AGENT_KNOWLEDGE_TYPE,
    VECTOR_TYPE,
)
from dataagent.vectorstore import RetrievalService


class FileStorage(Protocol):
    async def save(self, name: str, content: bytes) -> tuple[str, str]: ...

    async def delete(self, location: str) -> None: ...


def _safe_name(name: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]", "_", Path(name).name)
    return clean[:180] or "upload.bin"


async def read_upload_limited(upload: object, max_bytes: int) -> bytes:
    """Read an UploadFile without allowing an unbounded in-memory request."""

    if max_bytes <= 0:
        raise ConfigurationError("Upload size limit must be positive")
    read = getattr(upload, "read", None)
    if read is None:
        raise DataAgentError("Invalid upload")
    payload = bytearray()
    chunk_size = min(1024 * 1024, max_bytes + 1)
    while True:
        chunk = await read(chunk_size)
        if not chunk:
            break
        payload.extend(chunk)
        if len(payload) > max_bytes:
            raise DataAgentError(
                f"Upload exceeds the {max_bytes}-byte limit",
                status_code=413,
            )
    return bytes(payload)


class LocalFileStorage:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    async def save(self, name: str, content: bytes) -> tuple[str, str]:
        target_name = f"{uuid4().hex}_{_safe_name(name)}"
        target = await asyncio.to_thread(self._save_sync, target_name, content)
        return str(target), f"{self.settings.file_url_prefix.rstrip('/')}/{target_name}"

    async def delete(self, location: str) -> None:
        await asyncio.to_thread(self._delete_sync, location)

    def _save_sync(self, target_name: str, content: bytes) -> Path:
        target = self.settings.file_path.resolve() / target_name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)
        return target

    def _delete_sync(self, location: str) -> None:
        root = self.settings.file_path.resolve()
        target = Path(location).resolve()
        if target != root and root in target.parents:
            target.unlink(missing_ok=True)


class OssFileStorage:
    def __init__(self, settings: Settings) -> None:
        try:
            import oss2
        except ImportError as exc:
            raise ConfigurationError("Install the 'storage' extra to use OSS") from exc
        if not all(
            (
                settings.oss_endpoint,
                settings.oss_access_key_id,
                settings.oss_access_key_secret,
                settings.oss_bucket_name,
            )
        ):
            raise ConfigurationError("OSS credentials and bucket are required")
        auth = oss2.Auth(settings.oss_access_key_id, settings.oss_access_key_secret)
        self.bucket = oss2.Bucket(auth, settings.oss_endpoint, settings.oss_bucket_name)
        self.settings = settings

    async def save(self, name: str, content: bytes) -> tuple[str, str]:
        import asyncio

        key = f"{self.settings.file_path_prefix.strip('/')}/{uuid4().hex}_{_safe_name(name)}"
        await asyncio.to_thread(self.bucket.put_object, key, content)
        domain = self.settings.oss_custom_domain.rstrip("/")
        url = (
            f"{domain}/{key}"
            if domain
            else f"https://{self.settings.oss_bucket_name}.{self.settings.oss_endpoint}/{key}"
        )
        return key, url

    async def delete(self, location: str) -> None:
        import asyncio

        await asyncio.to_thread(self.bucket.delete_object, location)


def get_file_storage(settings: Settings | None = None) -> FileStorage:
    config = settings or get_settings()
    return OssFileStorage(config) if config.file_storage_type == "oss" else LocalFileStorage(config)


def storage_location_from_url(settings: Settings, value: str | None) -> str | None:
    """Recover the storage key from the public avatar URL returned by upload."""

    if not value:
        return None
    if settings.file_storage_type == "local":
        path = urlparse(value).path
        prefixes = (
            "/api/upload" + settings.file_url_prefix.rstrip("/") + "/",
            settings.file_url_prefix.rstrip("/") + "/",
        )
        for prefix in prefixes:
            if path.startswith(prefix):
                relative = path[len(prefix) :]
                if relative and Path(relative).name == relative:
                    return str(settings.file_path.resolve() / relative)
        candidate = Path(value)
        root = settings.file_path.resolve()
        resolved = candidate.resolve()
        return str(resolved) if resolved != root and root in resolved.parents else None

    if value.startswith(settings.file_path_prefix.strip("/") + "/"):
        return value
    expected_bases = [settings.oss_custom_domain.rstrip("/")]
    if settings.oss_bucket_name and settings.oss_endpoint:
        expected_bases.append(f"https://{settings.oss_bucket_name}.{settings.oss_endpoint}".rstrip("/"))
    if any(base and value.startswith(base + "/") for base in expected_bases):
        key = urlparse(value).path.lstrip("/")
        return key or None
    return None


class DocumentExtractor:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()

    async def extract(self, filename: str, content: bytes, content_type: str | None = None) -> str:
        suffix = Path(filename).suffix.lower()
        try:
            if suffix in {".txt", ".md", ".log", ".sql", ".py", ".html", ".xml"}:
                return content.decode("utf-8", errors="replace")
            if suffix == ".json":
                value = json.loads(content.decode("utf-8", errors="replace"))
                return json.dumps(value, ensure_ascii=False, indent=2)
            if suffix == ".jsonl":
                values = [
                    json.loads(line)
                    for line in content.decode("utf-8", errors="replace").splitlines()
                    if line.strip()
                ]
                return "\n".join(json.dumps(value, ensure_ascii=False) for value in values)
            if suffix in {".csv", ".tsv"}:
                dialect = "excel-tab" if suffix == ".tsv" else "excel"
                rows = csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")), dialect=dialect)
                return "\n".join(" | ".join(cell for cell in row) for row in rows)
            if suffix == ".pdf":
                reader = PdfReader(io.BytesIO(content))
                return "\n\n".join(page.extract_text() or "" for page in reader.pages)
            if suffix == ".docx":
                document = WordDocument(io.BytesIO(content))
                return "\n".join(paragraph.text for paragraph in document.paragraphs)
            if suffix in {".xlsx", ".xlsm"}:
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                lines: list[str] = []
                for sheet in workbook.worksheets:
                    lines.append(f"## {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        lines.append(" | ".join("" if value is None else str(value) for value in row))
                return "\n".join(lines)
        except Exception as exc:
            raise DataAgentError(f"Failed to parse {filename}: {exc}") from exc
        return await self._extract_with_tika(filename, content, content_type)

    async def _extract_with_tika(self, filename: str, content: bytes, content_type: str | None) -> str:
        try:
            headers = {
                "Accept": "text/plain",
                "Content-Disposition": f'attachment; filename="{_safe_name(filename)}"',
            }
            if content_type:
                headers["Content-Type"] = content_type
            async with httpx.AsyncClient(timeout=120) as client:
                response = await client.put(
                    f"{self.settings.tika_url.rstrip('/')}/tika", content=content, headers=headers
                )
            response.raise_for_status()
            return response.text
        except Exception as exc:
            if self.settings.tika_required:
                raise DataAgentError(f"Apache Tika extraction failed: {exc}", status_code=502) from exc
            suffix = Path(filename).suffix or "(none)"
            raise DataAgentError(
                f"Unsupported document type {suffix}; configure Tika for this format"
            ) from exc


def _window(values: list[str], size: int, overlap: int) -> list[str]:
    if not values:
        return []
    step = max(1, size - overlap)
    return [
        "".join(values[start : start + size]).strip()
        for start in range(0, len(values), step)
        if "".join(values[start : start + size]).strip()
    ]


class TextSplitter:
    def __init__(self, settings: Settings | None = None, llm: LLMClient | None = None) -> None:
        self.settings = settings or get_settings()
        self.llm = llm or get_llm_client()

    async def split(self, session: AsyncSession, text: str, splitter_type: str) -> list[str]:
        normalized = text.replace("\r\n", "\n").strip()
        if not normalized:
            return []
        splitter = splitter_type.lower()
        if splitter == "paragraph":
            chunks = self._paragraph(normalized)
        elif splitter == "sentence":
            chunks = self._sentence(normalized)
        elif splitter == "recursive":
            chunks = self._recursive(normalized)
        elif splitter == "semantic":
            chunks = await self._semantic(session, normalized)
        else:
            chunks = self._token(normalized)
        if len(chunks) > self.settings.text_max_chunks:
            raise DataAgentError(f"Document produced more than {self.settings.text_max_chunks} chunks")
        return chunks

    def _token(self, text: str) -> list[str]:
        try:
            import tiktoken

            encoding = tiktoken.get_encoding("cl100k_base")
            tokens = encoding.encode(text)
            size = self.settings.text_chunk_size
            overlap = self.settings.text_chunk_overlap
            step = max(1, size - overlap)
            return [
                encoding.decode(tokens[start : start + size]).strip() for start in range(0, len(tokens), step)
            ]
        except Exception:
            return _window(list(text), self.settings.text_chunk_size, self.settings.text_chunk_overlap)

    def _recursive(self, text: str) -> list[str]:
        separators = ["\n\n", "\n", "。", "！", "？", ". ", " "]
        pending = [text]
        chunks: list[str] = []
        for separator in separators:
            next_pending: list[str] = []
            for piece in pending:
                if len(piece) <= self.settings.text_chunk_size:
                    chunks.append(piece)
                else:
                    split = piece.split(separator)
                    next_pending.extend(part + separator for part in split if part.strip())
            pending = next_pending
        chunks.extend(
            part
            for piece in pending
            for part in _window(list(piece), self.settings.text_chunk_size, self.settings.text_chunk_overlap)
        )
        return self._merge_small(chunks)

    def _sentence(self, text: str) -> list[str]:
        sentences = [part for part in re.split(r"(?<=[。！？.!?])\s*", text) if part.strip()]
        chunks: list[str] = []
        current: list[str] = []
        for sentence in sentences:
            if current and len("".join(current)) + len(sentence) > self.settings.text_chunk_size:
                chunks.append("".join(current))
                current = (
                    current[-self.settings.text_sentence_overlap :]
                    if self.settings.text_sentence_overlap
                    else []
                )
            current.append(sentence)
        if current:
            chunks.append("".join(current))
        return self._merge_small(chunks)

    def _paragraph(self, text: str) -> list[str]:
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        chunks: list[str] = []
        current = ""
        for paragraph in paragraphs:
            candidate = f"{current}\n\n{paragraph}".strip()
            if current and len(candidate) > self.settings.text_chunk_size:
                chunks.append(current)
                current = current[-self.settings.text_paragraph_overlap_chars :] + "\n\n" + paragraph
            else:
                current = candidate
        if current:
            chunks.append(current)
        return self._merge_small(chunks)

    async def _semantic(self, session: AsyncSession, text: str) -> list[str]:
        sentences = [part for part in re.split(r"(?<=[。！？.!?])\s*", text) if part.strip()]
        if len(sentences) <= 1:
            return self._recursive(text)
        vectors = await self.llm.embed(session, sentences)
        chunks: list[str] = []
        current = sentences[0]
        for index in range(1, len(sentences)):
            left, right = vectors[index - 1], vectors[index]
            denom = sum(x * x for x in left) ** 0.5 * sum(x * x for x in right) ** 0.5
            similarity = sum(x * y for x, y in zip(left, right, strict=True)) / denom if denom else 0
            sentence = sentences[index]
            should_break = len(current) >= self.settings.text_semantic_min_chunk_size and (
                similarity < self.settings.text_semantic_threshold
                or len(current) + len(sentence) > self.settings.text_semantic_max_chunk_size
            )
            if should_break:
                chunks.append(current)
                current = sentence
            else:
                current += sentence
        chunks.append(current)
        return self._merge_small(chunks)

    def _merge_small(self, chunks: list[str]) -> list[str]:
        merged: list[str] = []
        for chunk in (item.strip() for item in chunks if item.strip()):
            if merged and len(chunk) < self.settings.text_min_chunk_chars:
                merged[-1] = f"{merged[-1]}\n{chunk}"
            else:
                merged.append(chunk)
        return merged


class KnowledgeIndexer:
    def __init__(
        self,
        retrieval: RetrievalService | None = None,
        splitter: TextSplitter | None = None,
    ) -> None:
        self.retrieval = retrieval or RetrievalService()
        self.splitter = splitter or TextSplitter()

    async def index_agent_knowledge(self, session: AsyncSession, knowledge_id: int) -> None:
        item = await session.get(AgentKnowledge, knowledge_id)
        if item is None or item.is_deleted:
            return
        item.embedding_status = "PROCESSING"
        item.error_msg = None
        await session.commit()
        try:
            knowledge_type = item.type.upper()
            if knowledge_type in {"QA", "FAQ"}:
                if not item.question or not item.content:
                    raise DataAgentError("Question and content are required for QA or FAQ knowledge")
                chunks = [item.question]
            elif knowledge_type == "DOCUMENT":
                chunks = await self.splitter.split(session, item.content or "", item.splitter_type)
            else:
                raise DataAgentError(f"Unsupported knowledge type: {item.type}")
            if not chunks:
                raise DataAgentError("No document content was extracted for vectorization")
            metadata = {
                AGENT_ID: item.agent_id,
                VECTOR_TYPE: AGENT_KNOWLEDGE,
                AGENT_KNOWLEDGE_ID: item.id,
                CONCRETE_AGENT_KNOWLEDGE_TYPE: knowledge_type,
            }
            await self.retrieval.replace(
                session,
                texts=chunks,
                metadata=metadata,
                old_filter={
                    AGENT_ID: item.agent_id,
                    VECTOR_TYPE: AGENT_KNOWLEDGE,
                    AGENT_KNOWLEDGE_ID: item.id,
                },
            )
            item.embedding_status = "COMPLETED"
        except Exception as exc:
            item.embedding_status = "FAILED"
            item.error_msg = str(exc)[:255]
        await session.commit()

    async def index_business_knowledge(self, session: AsyncSession, knowledge_id: int) -> None:
        item = await session.get(BusinessKnowledge, knowledge_id)
        if item is None or item.is_deleted:
            return
        item.embedding_status = "PROCESSING"
        item.error_msg = None
        await session.commit()
        try:
            content = "业务名词: {}, 说明: {}, 同义词: {}".format(
                item.business_term,
                item.description or "无",
                item.synonyms or "无",
            )
            metadata = {
                AGENT_ID: item.agent_id,
                VECTOR_TYPE: BUSINESS_TERM,
                BUSINESS_TERM_ID: item.id,
            }
            await self.retrieval.replace(
                session,
                texts=[content],
                metadata=metadata,
                old_filter={
                    AGENT_ID: item.agent_id,
                    VECTOR_TYPE: BUSINESS_TERM,
                    BUSINESS_TERM_ID: item.id,
                },
            )
            item.embedding_status = "COMPLETED"
        except Exception as exc:
            item.embedding_status = "FAILED"
            item.error_msg = str(exc)[:255]
        await session.commit()

    async def remove_agent_knowledge(self, item: AgentKnowledge) -> None:
        await self.retrieval.store.delete_by_metadata(
            {
                AGENT_ID: item.agent_id,
                VECTOR_TYPE: AGENT_KNOWLEDGE,
                AGENT_KNOWLEDGE_ID: item.id,
            }
        )

    async def remove_business_knowledge(self, item: BusinessKnowledge) -> None:
        await self.retrieval.store.delete_by_metadata(
            {
                AGENT_ID: item.agent_id,
                VECTOR_TYPE: BUSINESS_TERM,
                BUSINESS_TERM_ID: item.id,
            }
        )


async def recalled_evidence(
    session: AsyncSession, retrieval: RetrievalService, agent_id: int, query: str
) -> str:
    business_rows = (
        (
            await session.execute(
                select(BusinessKnowledge).where(
                    BusinessKnowledge.agent_id == agent_id,
                    BusinessKnowledge.is_deleted == 0,
                    BusinessKnowledge.is_recall == 1,
                )
            )
        )
        .scalars()
        .all()
    )
    agent_rows = (
        (
            await session.execute(
                select(AgentKnowledge).where(
                    AgentKnowledge.agent_id == agent_id,
                    AgentKnowledge.is_deleted == 0,
                    AgentKnowledge.is_recall == 1,
                )
            )
        )
        .scalars()
        .all()
    )

    filters: list[dict[str, object]] = []
    positions: dict[str, int] = {}
    if business_rows:
        positions["business"] = len(filters)
        filters.append(
            {
                AGENT_ID: agent_id,
                VECTOR_TYPE: BUSINESS_TERM,
                f"{BUSINESS_TERM_ID}In": [item.id for item in business_rows],
            }
        )
    if agent_rows:
        positions["agent"] = len(filters)
        filters.append(
            {
                AGENT_ID: agent_id,
                VECTOR_TYPE: AGENT_KNOWLEDGE,
                f"{AGENT_KNOWLEDGE_ID}In": [item.id for item in agent_rows],
            }
        )
    results = await retrieval.search_many(session, query, filter_sets=filters)
    business: list | BaseException = (
        results[positions["business"]] if "business" in positions else []
    )
    agent: list | BaseException = results[positions["agent"]] if "agent" in positions else []
    business_content = ""
    if not isinstance(business, BaseException):
        business_content = "\n".join(item.content for item in business)

    agent_content = ""
    if not isinstance(agent, BaseException):
        rows_by_id = {item.id: item for item in agent_rows}
        formatted: list[str] = []
        for index, document in enumerate(agent, start=1):
            try:
                item_id = int(document.metadata[AGENT_KNOWLEDGE_ID])
            except (KeyError, TypeError, ValueError):
                formatted.append(f"{index}. [来源: 知识库] {document.content}")
                continue
            item = rows_by_id.get(item_id)
            if item is None:
                continue
            if item.type.upper() in {"QA", "FAQ"}:
                formatted.append(
                    f"{index}. [来源: {item.title or '知识库'}] "
                    f"Q: {document.content} A: {item.content or ''}"
                )
            else:
                source = item.title or "文档"
                if item.source_filename:
                    source += f"-{item.source_filename}"
                formatted.append(f"{index}. [来源: {source}] {document.content}")
        agent_content = "\n".join(formatted)

    sections: list[str] = []
    if business_content:
        sections.append(render_prompt("business-knowledge", businessKnowledge=business_content))
    if agent_content:
        sections.append(render_prompt("agent-knowledge", agentKnowledge=agent_content))
    return "\n\n".join(sections) if sections else "无"
