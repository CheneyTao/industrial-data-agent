from __future__ import annotations

import io
from importlib.resources import files
from typing import Any

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import Response
from openpyxl import load_workbook
from sqlalchemy import delete, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from dataagent.database import get_session
from dataagent.documents import read_upload_limited
from dataagent.errors import DataAgentError, NotFoundError
from dataagent.models import AgentDatasource, SemanticModel
from dataagent.responses import ok
from dataagent.schemas import SemanticModelBatchImport, SemanticModelImportItem, SemanticModelPayload
from dataagent.serialization import serialize

router = APIRouter(prefix="/api/semantic-model", tags=["semantic-model"])


async def _item(session: AsyncSession, item_id: int) -> SemanticModel:
    value = await session.get(SemanticModel, item_id)
    if value is None:
        raise NotFoundError("Semantic model not found")
    return value


async def _datasource_id(session: AsyncSession, agent_id: int, supplied: int | None) -> int:
    statement = select(AgentDatasource).where(AgentDatasource.agent_id == agent_id)
    if supplied is not None:
        statement = statement.where(AgentDatasource.datasource_id == supplied)
    values = (
        await session.execute(
            statement.order_by(AgentDatasource.is_active.desc(), AgentDatasource.id.asc())
        )
    ).scalars().all()
    if not values:
        if supplied is not None:
            raise DataAgentError("Datasource is not associated with this agent")
        raise DataAgentError("Agent has no datasource")
    return int(values[0].datasource_id)


@router.get("")
async def list_semantic_models(
    keyword: str | None = None,
    agent_id: int | None = Query(None, alias="agentId"),
    session: AsyncSession = Depends(get_session),
):
    statement = select(SemanticModel)
    if keyword:
        like = f"%{keyword}%"
        statement = statement.where(
            or_(
                SemanticModel.table_name.ilike(like),
                SemanticModel.column_name.ilike(like),
                SemanticModel.business_name.ilike(like),
                SemanticModel.business_description.ilike(like),
            )
        )
    elif agent_id is not None:
        statement = statement.where(SemanticModel.agent_id == agent_id)
    rows = (await session.execute(statement.order_by(SemanticModel.id.desc()))).scalars()
    return ok([serialize(item) for item in rows], "success list semanticModel")


@router.get("/template/download")
async def download_template():
    path = files("dataagent").joinpath("assets/semantic_model_template.xlsx")
    if not path.is_file():
        raise NotFoundError("semantic_model_template.xlsx not found")
    return Response(
        path.read_bytes(),
        media_type="application/octet-stream",
        headers={"Content-Disposition": 'attachment; filename="semantic_model_template.xlsx"'},
    )


@router.get("/{item_id:int}")
async def get_semantic_model(item_id: int, session: AsyncSession = Depends(get_session)):
    return ok(serialize(await _item(session, item_id)), "success retrieve semanticModel")


@router.post("")
async def create_semantic_model(payload: SemanticModelPayload, session: AsyncSession = Depends(get_session)):
    data = payload.model_dump()
    # The upstream create DTO derives datasourceId from the agent association.
    data["datasource_id"] = await _datasource_id(session, payload.agent_id, None)
    data["status"] = 1
    value = SemanticModel(**data)
    session.add(value)
    await session.commit()
    return ok(True, "Semantic model created successfully")


@router.put("/{item_id:int}")
async def update_semantic_model(
    item_id: int, payload: SemanticModelPayload, session: AsyncSession = Depends(get_session)
):
    value = await _item(session, item_id)
    data = payload.model_dump()
    data["datasource_id"] = await _datasource_id(session, payload.agent_id, payload.datasource_id)
    for key, item in data.items():
        setattr(value, key, item)
    await session.commit()
    await session.refresh(value)
    return ok(serialize(value), "Semantic model updated successfully")


@router.delete("/batch")
async def batch_delete_semantic_models(ids: list[int], session: AsyncSession = Depends(get_session)):
    if not ids:
        raise DataAgentError("ID列表不能为空")
    await session.execute(delete(SemanticModel).where(SemanticModel.id.in_(ids)))
    await session.commit()
    return ok(True, "批量删除成功")


@router.put("/enable")
async def enable_semantic_models(ids: list[int], session: AsyncSession = Depends(get_session)):
    if not ids:
        raise DataAgentError("ID列表不能为空")
    await session.execute(update(SemanticModel).where(SemanticModel.id.in_(ids)).values(status=1))
    await session.commit()
    return ok(True, "Semantic models enabled successfully")


@router.put("/disable")
async def disable_semantic_models(ids: list[int], session: AsyncSession = Depends(get_session)):
    if not ids:
        raise DataAgentError("ID列表不能为空")
    await session.execute(update(SemanticModel).where(SemanticModel.id.in_(ids)).values(status=0))
    await session.commit()
    return ok(True, "Semantic models disabled successfully")


@router.delete("/{item_id:int}")
async def delete_semantic_model(item_id: int, session: AsyncSession = Depends(get_session)):
    value = await _item(session, item_id)
    await session.delete(value)
    await session.commit()
    return ok(True, "Semantic model deleted successfully")


async def _batch_import(
    session: AsyncSession, agent_id: int, items: list[SemanticModelImportItem]
) -> dict[str, Any]:
    datasource_id = await _datasource_id(session, agent_id, None)
    errors: list[str] = []
    success = 0
    for index, item in enumerate(items, start=1):
        try:
            existing = (
                await session.execute(
                    select(SemanticModel).where(
                        SemanticModel.agent_id == agent_id,
                        SemanticModel.table_name == item.table_name,
                        SemanticModel.column_name == item.column_name,
                    ).order_by(SemanticModel.id).limit(1)
                )
            ).scalar_one_or_none()
            data = item.model_dump(exclude={"status", "created_time"})
            if existing is None:
                if item.created_time is not None:
                    data["created_time"] = item.created_time
                session.add(
                    SemanticModel(
                        agent_id=agent_id,
                        datasource_id=datasource_id,
                        status=1,
                        **data,
                    )
                )
            else:
                for key, value in data.items():
                    setattr(existing, key, value)
            success += 1
        except Exception as exc:
            errors.append(f"row {index}: {exc}")
    await session.commit()
    return {"total": len(items), "successCount": success, "failCount": len(errors), "errors": errors}


@router.post("/batch-import")
async def batch_import_semantic_models(
    payload: SemanticModelBatchImport, session: AsyncSession = Depends(get_session)
):
    result = await _batch_import(session, payload.agent_id, payload.items)
    return ok(result, "批量导入完成")


@router.post("/import/excel")
async def import_semantic_excel(request: Request, session: AsyncSession = Depends(get_session)):
    form = await request.form()
    upload = form.get("file")
    try:
        agent_id = int(str(form.get("agentId")))
    except ValueError as exc:
        raise DataAgentError("agentId is required") from exc
    if upload is None or not hasattr(upload, "read"):
        raise DataAgentError("Excel file is required")
    content = await read_upload_limited(upload, request.app.state.settings.file_max_size)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise DataAgentError(f"Excel导入失败: {exc}") from exc
    items: list[SemanticModelImportItem] = []
    errors: list[str] = []
    for number, row in enumerate(rows[1:], start=2):
        values = list(row) + [None] * (8 - len(row))
        if not any(value not in (None, "") for value in values[:8]):
            continue
        try:
            items.append(
                SemanticModelImportItem(
                    table_name=str(values[0] or ""),
                    column_name=str(values[1] or ""),
                    business_name=str(values[2] or ""),
                    data_type=str(values[3] or ""),
                    synonyms=str(values[4] or "") or None,
                    business_description=str(values[5] or "") or None,
                    column_comment=str(values[6] or "") or None,
                    created_time=values[7] or None,
                )
            )
        except Exception as exc:
            errors.append(f"row {number}: {exc}")
    result = await _batch_import(session, agent_id, items)
    result["total"] += len(errors)
    result["failCount"] += len(errors)
    result["errors"] = errors + result["errors"]
    return ok(result, "Excel导入完成")
